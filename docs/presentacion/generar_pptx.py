# -*- coding: utf-8 -*-
"""Genera el PowerPoint de defensa de tesis CarDetector."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parent / "CarDetector_Defensa_Tesis.pptx"

# Paleta académica (azul oscuro + acento teal, sin púrpura)
NAVY = RGBColor(0x0F, 0x2C, 0x4C)
TEAL = RGBColor(0x1A, 0x7A, 0x7A)
SLATE = RGBColor(0x33, 0x3F, 0x4F)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_LINE = RGBColor(0x2E, 0xA3, 0xA3)


def _set_run(run, size=18, bold=False, color=SLATE, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_bg(slide, color=LIGHT):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _bar(slide, top=False):
    shape = slide.shapes.add_shape(
        1,  # rectangle
        Inches(0),
        Inches(0) if top else Inches(7.0),
        Inches(13.333),
        Inches(0.12) if top else Inches(0.5),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY if top else NAVY
    shape.line.fill.background()
    if not top:
        # footer accent strip already navy; add thin teal on top of footer
        accent = slide.shapes.add_shape(
            1, Inches(0), Inches(7.0), Inches(13.333), Inches(0.06)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = ACCENT_LINE
        accent.line.fill.background()


def _title_box(slide, text, left=0.6, top=0.35, width=12, height=0.7):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    _set_run(run, size=28, bold=True, color=NAVY)
    return box


def _bullets(slide, items, left=0.7, top=1.3, width=11.8, height=5.2, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = item.get("level", 0)
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = item["text"]
        _set_run(run, size=size - (2 if p.level else 0), bold=item.get("bold", False), color=SLATE)
    return box


def _notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.text = text.strip()


def _two_cols(slide, left_items, right_items, left_title, right_title):
    # left card title
    for col, title, items, x in [
        (0, left_title, left_items, 0.5),
        (1, right_title, right_items, 6.9),
    ]:
        t = slide.shapes.add_textbox(Inches(x), Inches(1.25), Inches(5.8), Inches(0.45))
        r = t.text_frame.paragraphs[0].add_run()
        r.text = title
        _set_run(r, size=20, bold=True, color=TEAL)
        _bullets(slide, items, left=x, top=1.85, width=5.8, height=4.5, size=16)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---------- 1 Portada ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s, NAVY)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.5), Inches(1.2))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = "CarDetector"
    _set_run(r, size=44, bold=True, color=WHITE)
    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(11.5), Inches(1.2))
    p = sub.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Sistema de censo vehicular y reconocimiento\nde matrículas argentinas en video"
    _set_run(r, size=22, color=RGBColor(0xC8, 0xE6, 0xE6))
    meta = s.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11.5), Inches(1.2))
    tf = meta.text_frame
    for i, line in enumerate(
        [
            "Trabajo Integrador Final — Ingeniería en Sistemas",
            "Lautaro Giménez Gil",
            "Defensa de tesis",
        ]
    ):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        _set_run(r, size=16, color=RGBColor(0xA8, 0xC0, 0xD8))
    _notes(
        s,
        """Buenos días / buenas tardes. Mi nombre es Lautaro Giménez Gil y presento CarDetector,
un sistema end-to-end para censar vehículos en video y reconocer matrículas argentinas.
La presentación dura unos 10–12 minutos: problema, solución técnica, resultados y demo.
Al final quedo a disposición de preguntas.""",
    )

    # ---------- 2 Agenda ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Agenda")
    _bullets(
        s,
        [
            {"text": "1. Problema y motivación"},
            {"text": "2. Objetivos y alcance"},
            {"text": "3. Arquitectura (frontend + backend)"},
            {"text": "4. Pipelines: censo y matrículas"},
            {"text": "5. Evaluación experimental y resultados"},
            {"text": "6. Limitaciones, conclusiones y trabajo futuro"},
            {"text": "7. Demostración en vivo (si el tiempo lo permite)"},
        ],
        size=20,
    )
    _notes(
        s,
        """Voy a recorrer primero el problema y los objetivos, después la arquitectura y los dos
pipelines de visión, luego los resultados del capítulo experimental, limitaciones y cierre.
Si hay tiempo, muestro una corrida corta en la interfaz Streamlit.""",
    )

    # ---------- 3 Problema ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Problema y motivación")
    _bullets(
        s,
        [
            {"text": "El conteo manual de tránsito es costoso, lento y propenso a error."},
            {"text": "Hace falta automatizar:"},
            {"text": "Censo por tipo de vehículo (auto, moto, bus, camión)", "level": 1},
            {"text": "Lectura de patentes argentinas (formatos 6 y 7 caracteres)", "level": 1},
            {"text": "Condiciones reales: día/noche, distintas resoluciones y ángulos"},
            {"text": "Aporte: un sistema usable (UI + API) reproducible, no solo un notebook."},
        ],
        size=19,
    )
    _notes(
        s,
        """El problema de partida es operativo: censar y leer patentes a mano no escala.
La tesis no se queda en un script aislado: entrega un producto técnico mínimo —
frontend para operadores y backend con pipelines de visión — pensado para demo y evaluación.""",
    )

    # ---------- 4 Objetivos ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Objetivos")
    _two_cols(
        s,
        [
            {"text": "Diseñar e implementar un sistema de censo vehicular en video."},
            {"text": "Detectar y leer matrículas argentinas (OCR)."},
            {"text": "Exponer ambos flujos vía API e interfaz web."},
            {"text": "Evaluar el sistema con un protocolo experimental documentado."},
        ],
        [
            {"text": "Entrenar desde cero un detector SOTA."},
            {"text": "Sustituir sistemas comerciales de peaje."},
            {"text": "Tracking multi-cámara en tiempo real de producción."},
            {"text": "Dataset público masivo versionado (queda como futuro)."},
        ],
        "Objetivos específicos",
        "Fuera de alcance (delimitación)",
    )
    _notes(
        s,
        """Los objetivos específicos son: censo, OCR de patentes AR, integración API+UI y evaluación.
Deliberadamente queda fuera entrenar un detector desde cero o competir con peajes comerciales:
reutilizo modelos preentrenados adaptados al dominio argentino y me concentro en la integración.""",
    )

    # ---------- 5 Arquitectura ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Arquitectura del sistema")
    _bullets(
        s,
        [
            {"text": "Frontend (Streamlit): carga de video, censo / matrículas, tablas y CSV."},
            {"text": "Backend (Flask): API REST, jobs asíncronos, carga única de modelos."},
            {"text": "Pipeline A — Censo: YOLOv5s (COCO) + tracking por centroides."},
            {"text": "Pipeline B — Patentes: YOLOv4-tiny (JustAnotherAlpr) + OCR ConvALPR."},
            {"text": "Comunicación: HTTP multipart → JSON (conteos, placas, timestamps)."},
            {"text": "Diseño headless: sin GUI en servidor; preview opcional en desktop."},
        ],
        size=18,
    )
    _notes(
        s,
        """Arquitectónicamente hay dos repos: cardetector-front y cardetector-backend.
El usuario sube un video en Streamlit; el backend corre el pipeline pedido y devuelve JSON.
Separé censo y patentes porque son problemas distintos: conteo multi-clase vs. localización + OCR.
Los modelos se cargan una vez al arrancar para no pagar el costo en cada request.""",
    )

    # ---------- 6 Censo ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Pipeline de censo vehicular")
    _bullets(
        s,
        [
            {"text": "Detector: YOLOv5s preentrenado en COCO (PyTorch / torch.hub)."},
            {"text": "Clases: auto, moto, bus, camión (IDs COCO 2, 3, 5, 7)."},
            {"text": "Tracking: asociación por centroides entre frames (ID estable en la escena)."},
            {"text": "Filtros anti-falsos positivos: confianza, área y aspecto (p. ej. carteles)."},
            {"text": "Salida: total por clase + total de vehículos únicos en el video."},
            {"text": "Optimización: stride de frames y downscale para fluidez en CPU."},
        ],
        size=18,
    )
    _notes(
        s,
        """Para el censo uso YOLOv5s de COCO: es un estándar conocido y suficiente para las cuatro clases.
El tracking por centroides evita contar el mismo auto muchas veces al pasar frame a frame.
Agregué filtros porque en calle hay carteles o formas que el modelo confunde con bus/camión.
Como corro en CPU (GPU AMD sin CUDA en PyTorch), uso stride y downscale para que sea usable.""",
    )

    # ---------- 7 Patentes ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Pipeline de matrículas (Argentina)")
    _bullets(
        s,
        [
            {"text": "Detección de placa: YOLOv4-tiny (JustAnotherAlpr), dominio AR."},
            {"text": "OCR: ConvALPR (modelo CNN para formato argentino)."},
            {"text": "Soporta condición día/noche como parámetro del pipeline."},
            {"text": "Salida: placas únicas + lecturas crudas / timestamps para análisis."},
            {"text": "Alternativas configurables (TF SavedModel + Keras.h5) vía .env."},
            {"text": "Enfoque: reutilizar pesos especializados, no reentrenar desde cero."},
        ],
        size=18,
    )
    _notes(
        s,
        """Las patentes argentinas tienen un formato particular; por eso no usé un OCR genérico.
JustAnotherAlpr aporta el detector de placa y ConvALPR el reconocimiento de caracteres.
El sistema admite día/noche porque la iluminación cambia mucho el contraste de la placa.
Documenté backends alternativos, pero la configuración por defecto de la tesis es darknet_ar + convalpr.""",
    )

    # ---------- 8 Frontend ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Interfaz de usuario (Streamlit)")
    _bullets(
        s,
        [
            {"text": "Pensada para operadores no técnicos: subir video y elegir flujo."},
            {"text": "Modos: Censo / Matrículas (con selector día-noche)."},
            {"text": "Feedback de progreso vía jobs + polling (preview JPEG opcional)."},
            {"text": "Resultados en tabla y descarga CSV para análisis posterior."},
            {"text": "Cliente HTTP desacoplado del backend (API_URL configurable)."},
        ],
        size=18,
    )
    _notes(
        s,
        """El frontend no es un adorno: es la cara usable del aporte.
Streamlit me permitió iterar rápido: upload, botones de censo/patentes, tablas y CSV.
El polling a jobs asíncronos evita que la UI se congele en videos largos.""",
    )

    # ---------- 9 Metodología ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Metodología experimental")
    _bullets(
        s,
        [
            {"text": "Corpus de videos locales (día/noche, distintas resoluciones)."},
            {"text": "Script batch: scripts/run_experiments.py → CSV + predicciones."},
            {"text": "Censo: comparar total predicho vs. conteo de referencia (error %)."},
            {"text": "OCR: precisión / recall / F1 sobre placas (con revisión manual)."},
            {"text": "Honestidad metodológica: distinguir GT manual vs. estimaciones."},
            {"text": "Ejemplo con GT manual: test1.mp4 (censo 5/5, F1 OCR ≈ 0.75)."},
        ],
        size=18,
    )
    _notes(
        s,
        """Para el capítulo de resultados automaticé corridas sobre una docena de videos.
Las métricas serias requieren ground truth: en test1 hice revisión manual.
En otros videos usé estimaciones solo para orientar la presentación; lo aclaro en la defensa
para no vender números sin etiquetado completo como si fueran GT cerrado.""",
    )

    # ---------- 10 Resultados ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Resultados (síntesis)")
    _bullets(
        s,
        [
            {"text": "12 videos procesados en batch (promedio ~22 s de duración)."},
            {"text": "test1.mp4 (GT manual): error de censo 0% (5/5) · F1 OCR ≈ 0.75"},
            {"text": "Escenas más densas / lejanas: mayor error de censo y OCR incompleto."},
            {"text": "Noche (ba_noche_720): el sistema produce lecturas, con más ruido OCR."},
            {"text": "Hallazgo: la integración funciona; el cuello de botella es calidad de placa"},
            {"text": "y tracking simple ante oclusiones / multi-objeto denso.", "level": 1},
        ],
        size=17,
    )
    _notes(
        s,
        """En test1, con ground truth manual, el censo acertó 5 de 5 y el OCR tuvo F1 alrededor de 0.75:
es el resultado más sólido para mostrar.
En videos densos o de noche el sistema sigue corriendo, pero el OCR se degrada:
placas chicas, motion blur, iluminación. Eso alimenta las limitaciones y el trabajo futuro.""",
    )

    # ---------- 11 Limitaciones ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Limitaciones")
    _bullets(
        s,
        [
            {"text": "Tracking por centroides (no SORT/ByteTrack): falla en escenas saturadas."},
            {"text": "OCR orientado a formato fijo de placa del dominio de entrenamiento."},
            {"text": "Jobs en memoria: se pierden al reiniciar el proceso."},
            {"text": "Ejecución en CPU: latencia alta en 1080p / 60 fps sin stride."},
            {"text": "GT completo solo en parte del corpus (resto estimado / pendiente)."},
            {"text": "Pesos fuera de Git (tamaño); hay que montarlos en despliegue."},
        ],
        size=18,
    )
    _notes(
        s,
        """Prefiero explicitar limitaciones: demuestra criterio de ingeniería.
El tracking simple es suficiente para la tesis pero no para producción densa.
Los jobs in-memory alcanzan para demo; en producción iría a Redis/RQ.
Y la evaluación completa con GT en todo el corpus sigue siendo trabajo abierto.""",
    )

    # ---------- 12 Conclusiones ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Conclusiones y trabajo futuro")
    _two_cols(
        s,
        [
            {"text": "Sistema end-to-end operativo: UI + API + dos pipelines."},
            {"text": "Censo viable con YOLOv5 + centroides en escenas moderadas."},
            {"text": "Patentes AR con modelos de dominio especializado."},
            {"text": "Protocolo experimental reproducible documentado."},
        ],
        [
            {"text": "Tracking SORT/ByteTrack + métricas MOT."},
            {"text": "Cola persistente (Redis) y despliegue Docker endurecido."},
            {"text": "Dataset versionado + más GT manual."},
            {"text": "Unificar stacks ML / aceleración GPU nativa."},
        ],
        "Conclusiones",
        "Trabajo futuro",
    )
    _notes(
        s,
        """Cierro destacando el aporte: no un paper de un modelo nuevo, sino un sistema integrado
evaluable. El futuro claro es mejor tracking, más ground truth y despliegue más robusto.""",
    )

    # ---------- 13 Demo ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s)
    _bar(s, top=True)
    _title_box(s, "Demostración")
    _bullets(
        s,
        [
            {"text": "1. Backend: python app.py → health OK"},
            {"text": "2. Frontend: streamlit run main.py"},
            {"text": "3. Subir video corto (p. ej. test1.mp4)"},
            {"text": "4. Ejecutar Censo → mostrar totales / CSV"},
            {"text": "5. Ejecutar Matrículas → mostrar placas únicas"},
            {"text": "Plan B: video pregrabado / capturas si falla el live."},
        ],
        size=20,
    )
    _notes(
        s,
        """Para la demo: backend y frontend ya levantados antes de empezar.
Uso un video corto y conocido (test1) para no improvisar tiempos.
Si algo falla, paso a capturas o al video resultado ya generado — nunca pelear con el live.""",
    )

    # ---------- 14 Gracias ----------
    s = prs.slides.add_slide(blank)
    _add_bg(s, NAVY)
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.0))
    r = box.text_frame.paragraphs[0].add_run()
    r.text = "Gracias"
    _set_run(r, size=48, bold=True, color=WHITE)
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(11.5), Inches(1.0))
    p = sub.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "¿Preguntas?"
    _set_run(r, size=28, color=RGBColor(0xC8, 0xE6, 0xE6))
    foot = s.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(11.5), Inches(0.8))
    p = foot.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Lautaro Giménez Gil — CarDetector"
    _set_run(r, size=16, color=RGBColor(0xA8, 0xC0, 0xD8))
    _notes(
        s,
        """Gracias. Quedo a disposición de las preguntas del tribunal.
Si preguntan por GPU: trabajo en CPU por limitación del stack CUDA/AMD; por eso stride/downscale.
Si preguntan por métricas: priorizo test1 con GT manual y soy transparente con el resto.""",
    )

    prs.save(OUT)
    print(f"OK -> {OUT}")


if __name__ == "__main__":
    build()
